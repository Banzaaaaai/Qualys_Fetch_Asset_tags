"""Fetch all Qualys asset tags (EU2 platform) and export to CSV + XLSX.

Credentials are never read from argv/hardcoded; they come from the
QUALYS_USERNAME / QUALYS_PASSWORD env vars, falling back to
qualys_creds.txt next to this script (format: "key:\\tvalue" per line).
"""

import csv
import datetime
import os
import sys
import time
import xml.etree.ElementTree as ET

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BASE_URL = "https://qualysapi.qg2.apps.qualys.eu"
TAG_SEARCH_URL = f"{BASE_URL}/qps/rest/2.0/search/am/tag"

INITIAL_DELAY = 1
MAX_RETRIES = 5
REQUEST_TIMEOUT = 60
MAX_RATE_LIMIT_SLEEP = 300
TAG_SEARCH_PAGE_SIZE = 100

HEADERS = {
    "Content-Type": "text/xml",
    "X-Requested-With": "Claude Code",
}

COLUMNS = [
    "tag_id",
    "tag_name",
    "parent_tag_id",
    "parent_tag_name",
    "criticality",
    "tag_type",
    "rule_type",
    "rule_value",
    "color_hex",
    "color_meaning",
]

# CSB003 section 3.2 color scheme. Multiple representative shades per
# meaning; unknown/absent colors, or colors too far from every shade,
# resolve to "" (best-effort match).
COLOR_SCHEME = {
    "Business-critical": ["FF0000", "DC143C", "B22222", "8B0000", "E32636"],
    "Negative signals": ["FFA500", "FF8C00", "FF7F50", "FF6347"],
    "Users/roles/departments": ["FFFF00", "FFD700", "EEE8AA", "F0E68C"],
    "Successful/positive": ["008000", "00FF00", "228B22", "32CD32", "006400"],
    "IPs/domains/asset groups": ["00008B", "000080", "191970", "0000CD", "00468B"],
    "Technologies": ["800080", "9932CC", "8A2BE2", "6A0DAD", "4B0082"],
    "Projects": ["FF00FF", "FF1493", "C71585", "DA70D6"],
    "Vulnerability QID/EOL": ["008080", "20B2AA", "008B8B", "5F9EA0"],
}
COLOR_MATCH_THRESHOLD = 120.0  # max RGB euclidean distance to accept a match


def load_credentials():
    username = os.environ.get("QUALYS_USERNAME")
    password = os.environ.get("QUALYS_PASSWORD")
    if username and password:
        return username, password

    creds_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "qualys_creds.txt")
    if not os.path.isfile(creds_path):
        raise RuntimeError(
            "No credentials found: set QUALYS_USERNAME/QUALYS_PASSWORD or provide qualys_creds.txt"
        )

    creds = {}
    with open(creds_path, "r", encoding="utf-8") as f:
        for line in f:
            if ":" not in line:
                continue
            key, _, value = line.partition(":")
            creds[key.strip().lower()] = value.strip()

    username = creds.get("user")
    password = creds.get("pass")
    if not username or not password:
        raise RuntimeError("qualys_creds.txt is missing 'user' and/or 'pass' entries")
    return username, password


def local_tag(tag):
    """Strip a namespace prefix like {ns}tag -> tag."""
    return tag.split("}", 1)[-1] if "}" in tag else tag


def find_child(elem, name):
    for child in elem:
        if local_tag(child.tag) == name:
            return child
    return None


def get_text(elem, name):
    child = find_child(elem, name)
    if child is None or child.text is None:
        return ""
    return child.text.strip()


def hex_to_rgb(hex_str):
    h = hex_str.strip().lstrip("#")
    if len(h) != 6:
        return None
    try:
        return tuple(int(h[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return None


def classify_color(color_hex):
    if not color_hex:
        return ""
    rgb = hex_to_rgb(color_hex)
    if rgb is None:
        return ""

    best_meaning = ""
    best_dist = None
    for meaning, shades in COLOR_SCHEME.items():
        for shade in shades:
            shade_rgb = hex_to_rgb(shade)
            dist = sum((a - b) ** 2 for a, b in zip(rgb, shade_rgb)) ** 0.5
            if best_dist is None or dist < best_dist:
                best_dist = dist
                best_meaning = meaning

    if best_dist is not None and best_dist <= COLOR_MATCH_THRESHOLD:
        return best_meaning
    return ""


def build_request_xml(last_id, page_size):
    return (
        "<ServiceRequest>"
        "<preferences>"
        f"<limitResults>{page_size}</limitResults>"
        "</preferences>"
        "<filters>"
        f'<Criteria field="id" operator="GREATER">{last_id}</Criteria>'
        "</filters>"
        "</ServiceRequest>"
    )


def post_with_retry(session, xml_body, page_num):
    delay = INITIAL_DELAY
    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = session.post(
                TAG_SEARCH_URL,
                data=xml_body.encode("utf-8"),
                headers=HEADERS,
                timeout=REQUEST_TIMEOUT,
            )
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as exc:
            last_error = str(exc)
            wait = min(delay, MAX_RATE_LIMIT_SLEEP)
            if attempt == MAX_RETRIES:
                break
            print(
                f"  [page {page_num}] network error ({exc}); retrying in {wait}s "
                f"(attempt {attempt}/{MAX_RETRIES})"
            )
            time.sleep(wait)
            delay = min(delay * 2, MAX_RATE_LIMIT_SLEEP)
            continue

        if resp.status_code == 200:
            root = ET.fromstring(resp.text)
            response_code = get_text(root, "responseCode")
            if response_code and response_code != "SUCCESS":
                error_msg = get_text(root, "responseErrorDetails") or resp.text[:500]
                raise RuntimeError(
                    f"Page {page_num}: Qualys API returned {response_code}: {error_msg}"
                )
            return root

        if resp.status_code in (409, 429):
            wait_header = resp.headers.get("X-RateLimit-ToWait-Sec") or resp.headers.get(
                "Retry-After"
            )
            try:
                wait = min(float(wait_header), MAX_RATE_LIMIT_SLEEP) if wait_header else min(
                    delay, MAX_RATE_LIMIT_SLEEP
                )
            except ValueError:
                wait = min(delay, MAX_RATE_LIMIT_SLEEP)

            last_error = f"HTTP {resp.status_code} (rate-limited)"
            if attempt == MAX_RETRIES:
                break
            print(
                f"  [page {page_num}] rate-limited (HTTP {resp.status_code}); waiting {wait}s "
                f"(attempt {attempt}/{MAX_RETRIES})"
            )
            time.sleep(wait)
            delay = min(delay * 2, MAX_RATE_LIMIT_SLEEP)
            continue

        last_error = f"HTTP {resp.status_code}: {resp.text[:500]}"
        wait = min(delay, MAX_RATE_LIMIT_SLEEP)
        if attempt == MAX_RETRIES:
            break
        print(
            f"  [page {page_num}] {last_error}; retrying in {wait}s "
            f"(attempt {attempt}/{MAX_RETRIES})"
        )
        time.sleep(wait)
        delay = min(delay * 2, MAX_RATE_LIMIT_SLEEP)

    raise RuntimeError(f"Page {page_num}: giving up after {MAX_RETRIES} attempts: {last_error}")


def parse_tag_elements(data_elem):
    tags = []
    if data_elem is None:
        return tags
    for tag_elem in data_elem:
        if local_tag(tag_elem.tag) != "Tag":
            continue
        rule_type = get_text(tag_elem, "ruleType")
        is_static = rule_type in ("", "STATIC")
        tags.append(
            {
                "tag_id": get_text(tag_elem, "id"),
                "tag_name": get_text(tag_elem, "name"),
                "parent_tag_id": get_text(tag_elem, "parentTagId"),
                "criticality": get_text(tag_elem, "criticalityScore"),
                "tag_type": "Static" if is_static else "Dynamic",
                "rule_type": rule_type,
                "rule_value": "" if is_static else get_text(tag_elem, "ruleText"),
                "color_hex": get_text(tag_elem, "color"),
            }
        )
    return tags


def fetch_all_tags(session):
    all_tags = []
    last_id = 0
    page_num = 0

    while True:
        page_num += 1
        xml_body = build_request_xml(last_id, TAG_SEARCH_PAGE_SIZE)
        root = post_with_retry(session, xml_body, page_num)

        data_elem = find_child(root, "data")
        page_tags = parse_tag_elements(data_elem)
        all_tags.extend(page_tags)

        has_more = get_text(root, "hasMoreRecords").lower() == "true"
        last_id_text = get_text(root, "lastId")
        if last_id_text:
            last_id = int(last_id_text)
        else:
            numeric_ids = [int(t["tag_id"]) for t in page_tags if t["tag_id"].isdigit()]
            if numeric_ids:
                last_id = max(last_id, max(numeric_ids))

        print(f"  [page {page_num}] fetched {len(page_tags)} tags (last_id={last_id})")

        if not has_more or not page_tags:
            break
        time.sleep(INITIAL_DELAY)

    return all_tags


def resolve_parent_names(tags):
    id_to_name = {t["tag_id"]: t["tag_name"] for t in tags if t["tag_id"]}
    unresolved = []
    for t in tags:
        parent_id = t["parent_tag_id"]
        if not parent_id:
            t["parent_tag_name"] = ""
            continue
        name = id_to_name.get(parent_id)
        if name is None:
            unresolved.append(t["tag_id"])
            t["parent_tag_name"] = ""
        else:
            t["parent_tag_name"] = name
    return unresolved


def finalize_rows(tags):
    for t in tags:
        t["color_meaning"] = classify_color(t["color_hex"])
    tags.sort(key=lambda t: (t["parent_tag_name"], t["tag_name"]))
    return tags


def write_csv(tags, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        for t in tags:
            writer.writerow({col: t.get(col, "") for col in COLUMNS})


def write_xlsx(tags, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tags"

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    rule_value_col = COLUMNS.index("rule_value") + 1
    for t in tags:
        ws.append([t.get(col, "") for col in COLUMNS])

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=rule_value_col).alignment = Alignment(
            wrap_text=True, vertical="top"
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    widths = {
        "tag_id": 10,
        "tag_name": 30,
        "parent_tag_id": 14,
        "parent_tag_name": 30,
        "criticality": 12,
        "tag_type": 10,
        "rule_type": 20,
        "rule_value": 60,
        "color_hex": 10,
        "color_meaning": 24,
    }
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 15)

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    ws2.append(["Total tags", len(tags)])
    ws2.append([])

    ws2.append(["Tag Type", "Count"])
    ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
    ws2.cell(row=ws2.max_row, column=2).font = Font(bold=True)
    type_counts = {}
    for t in tags:
        type_counts[t["tag_type"]] = type_counts.get(t["tag_type"], 0) + 1
    for tag_type in sorted(type_counts):
        ws2.append([tag_type, type_counts[tag_type]])
    ws2.append([])

    ws2.append(["Rule Type", "Count"])
    ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
    ws2.cell(row=ws2.max_row, column=2).font = Font(bold=True)
    rule_counts = {}
    for t in tags:
        key = t["rule_type"] or "(none)"
        rule_counts[key] = rule_counts.get(key, 0) + 1
    for rule_type in sorted(rule_counts):
        ws2.append([rule_type, rule_counts[rule_type]])

    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 12

    wb.save(path)


def print_summary(tags, unresolved_parent_ids):
    total = len(tags)
    static_count = sum(1 for t in tags if t["tag_type"] == "Static")
    dynamic_count = sum(1 for t in tags if t["tag_type"] == "Dynamic")
    no_color = sum(1 for t in tags if not t["color_hex"])
    no_parent = sum(1 for t in tags if not t["parent_tag_id"])

    print()
    print("=== Qualys Asset Tags Summary ===")
    print(f"Total tags:        {total}")
    print(f"Static tags:       {static_count}")
    print(f"Dynamic tags:      {dynamic_count}")
    print(f"No color set:      {no_color}")
    print(f"No parent (root):  {no_parent}")
    if unresolved_parent_ids:
        print(f"Unresolved parent tag_ids ({len(unresolved_parent_ids)}): {', '.join(unresolved_parent_ids)}")
    else:
        print("Unresolved parent tag_ids: none")


def main():
    username, password = load_credentials()

    session = requests.Session()
    session.auth = (username, password)

    print(f"Fetching all asset tags from {TAG_SEARCH_URL} ...")
    tags = fetch_all_tags(session)
    print(f"Fetched {len(tags)} tags total across all pages.")

    unresolved = resolve_parent_names(tags)
    tags = finalize_rows(tags)

    today = datetime.date.today().isoformat()
    csv_path = f"qualys_tags_{today}.csv"
    xlsx_path = f"qualys_tags_{today}.xlsx"

    write_csv(tags, csv_path)
    write_xlsx(tags, xlsx_path)

    print(f"Wrote {csv_path}")
    print(f"Wrote {xlsx_path}")

    print_summary(tags, unresolved)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
